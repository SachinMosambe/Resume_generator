"""Synthetic resume fixtures. Every person, company, and contact detail is fictional.

Each fixture has ground truth in ``EXPECTED`` so tests and ``scripts/fidelity_report.py``
can measure extraction, role association, and final-output coverage separately.
"""
from __future__ import annotations

from pathlib import Path

import pymupdf as fitz
from docx import Document

PERSON = "Jordan Rivera"
EMAIL = "jordan.rivera@example.com"
PHONE = "(555) 010-2030"

PROMOTION_ROLES = [
    ("Acme Corp", "Software Engineer", "Jan 2016 - Dec 2018", ["Built the billing service in Python for retail clients."]),
    ("Acme Corp", "Senior Software Engineer", "Jan 2019 - Jun 2021", []),
    ("Acme Corp", "Staff Engineer", "Jul 2021 - Present", ["Led the migration to event-driven payments across teams."]),
]

TWO_COL_ROLES = [
    ("Globex", "Data Engineer", "Mar 2019 - Present", ["Built streaming ETL pipelines for logistics data."]),
    ("Initech", "Analyst", "Jan 2015 - Feb 2019", ["Automated weekly finance reports for leadership."]),
]

LONG_ROLES = [
    (f"Company{i} Ltd", f"Engineer Level {i}", f"Jan {2000 + i} - Dec {2000 + i}", 30 if i == 0 else 4)
    for i in range(12)
]


def _long_bullet(i: int, j: int) -> str:
    return f"Delivered measurable outcome {i}-{j} for the platform team with clear ownership."


EXPECTED: dict[str, dict] = {
    "single_column_pdf": {
        "roles": [(c, t, d) for c, t, d, _ in PROMOTION_ROLES],
        "facts": [b for *_, bs in PROMOTION_ROLES for b in bs] + ["B.S. Computer Science", "State University"],
        "contact": {"email": EMAIL},
    },
    "two_column_pdf": {
        "roles": [(c, t, d) for c, t, d, _ in TWO_COL_ROLES],
        "facts": [b for *_, bs in TWO_COL_ROLES for b in bs]
        + ["Python", "Go", "SQL", "B.S. Physics", "Tech University", "Streaming joins at scale"],
        "sections": ["Publications"],
        "contact": {"email": EMAIL},
    },
    "sidebar_pdf": {
        "roles": [(c, t, d) for c, t, d, _ in TWO_COL_ROLES],
        "facts": [b for *_, bs in TWO_COL_ROLES for b in bs]
        + ["Python", "Spark", "Spanish", "AWS Certified Data Engineer - Associate",
           "Data engineer with ten years building analytics platforms for retail."],
    },
    "scanned_pdf": {
        "roles": [(c, t, d) for c, t, d, _ in PROMOTION_ROLES],
        "facts": [b for *_, bs in PROMOTION_ROLES for b in bs],
        "needs_ocr": True,
    },
    "mixed_pdf": {
        "roles": [(c, t, d) for c, t, d, _ in TWO_COL_ROLES],
        "facts": ["Built streaming ETL pipelines for logistics data.", "B.S. Physics", "Certified Scrum Master 2020"],
        "needs_ocr": True,
    },
    "docx_tables": {
        "roles": [("Globex", "Data Engineer", "Mar 2019 - Present")],
        "facts": ["Built streaming ETL pipelines for logistics data.", "B.S. Physics", "Tech University"],
    },
    "docx_header_contact": {
        "roles": [("Initech", "Analyst", "Jan 2015 - Feb 2019")],
        "facts": ["Automated weekly finance reports for leadership."],
        "contact": {"email": EMAIL, "name": PERSON},
    },
    "promotions": {
        "roles": [(c, t, d) for c, t, d, _ in PROMOTION_ROLES],
        "facts": [b for *_, bs in PROMOTION_ROLES for b in bs],
    },
    "overlapping": {
        "roles": [
            ("Globex", "Data Engineer", "Mar 2019 - Present"),
            ("Umbrella Labs", "Part-time Instructor", "Jan 2019 - Dec 2020"),
            ("Initech", "Analyst", "Jan 2015 - Feb 2019"),
        ],
        "facts": [
            "Built streaming ETL pipelines for logistics data.",
            "Taught evening data engineering classes to adult learners.",
            "Automated weekly finance reports for leadership.",
        ],
    },
    "odd_headings": {
        "roles": [("Globex", "Data Engineer", "Mar 2019 - Present")],
        "facts": ["Built streaming ETL pipelines for logistics data.", "Kubernetes", "Terraform"],
    },
    "long_resume": {
        "roles": [(c, t, d) for c, t, d, _ in LONG_ROLES],
        "facts": [_long_bullet(i, j) for i, (*_, n) in enumerate(LONG_ROLES) for j in range(n)],
    },
    "references": {
        "roles": [("Initech", "Analyst", "Jan 2015 - Feb 2019")],
        "facts": ["Automated weekly finance reports for leadership."],
        "contact": {"email": EMAIL},
        "not_candidate": {"email": "pat.manager@example.org", "phone_digits": "0998877"},
        "not_in_output": ["pat.manager@example.org", "099-8877"],
    },
    "unknown_sections": {
        "roles": [("Initech", "Analyst", "Jan 2015 - Feb 2019")],
        "facts": [
            "Automated weekly finance reports for leadership.",
            "Rivera J. Streaming joins at scale. 2022.",
            "Mentor, Code Club, 2020 - Present",
            "US 0,000,001 Event replay system",
        ],
        "sections": ["Publications", "Volunteering", "Patents"],
    },
}

TEMPLATE_SAMPLE_LINES = ["Sample Candidate Name", "Sample person with 10 years at SampleCo.", "Lead, SampleCo, 2010 - 2020"]


def _pdf_lines(path: Path, pages: list[list[tuple[float, float, str, float]]]) -> Path:
    doc = fitz.open()
    for items in pages:
        page = doc.new_page(width=612, height=792)
        for x, y, text, size in items:
            page.insert_text((x, y), text, fontsize=size, fontname="helv")
    doc.save(path)
    doc.close()
    return path


def _single_column_items() -> list[tuple[float, float, str, float]]:
    y = 60.0
    items = [(50.0, y, PERSON, 16.0)]
    for text in [EMAIL, PHONE, "", "EXPERIENCE"]:
        y += 18
        if text:
            items.append((50.0, y, text, 10.0))
    for company, title, dates, bullets in PROMOTION_ROLES:
        y += 18
        items.append((50.0, y, f"{title} | {company} | {dates}", 10.0))
        for b in bullets:
            y += 14
            items.append((60.0, y, f"- {b}", 10.0))
    y += 24
    items += [(50.0, y, "EDUCATION", 10.0), (50.0, y + 14, "B.S. Computer Science, State University, 2015", 10.0)]
    return items


def single_column_pdf(out: Path) -> Path:
    return _pdf_lines(out / "single_column.pdf", [_single_column_items()])


def _two_column_items() -> list[tuple[float, float, str, float]]:
    full = [(50.0, 50.0, PERSON, 16.0), (50.0, 68.0, f"{EMAIL} | {PHONE}", 10.0)]
    left = [
        (50.0, 110.0, "SKILLS", 10.0), (50.0, 124.0, "Python, Go, SQL", 10.0),
        (50.0, 160.0, "EDUCATION", 10.0), (50.0, 174.0, "B.S. Physics", 10.0), (50.0, 188.0, "Tech University 2014", 10.0),
    ]
    right = [(250.0, 110.0, "EXPERIENCE", 10.0)]
    y = 110.0
    for company, title, dates, bullets in TWO_COL_ROLES:
        y += 18
        right.append((250.0, y, f"{title} | {company} | {dates}", 10.0))
        for b in bullets:
            y += 14
            right.append((260.0, y, f"- {b}", 10.0))
    bottom = [
        (50.0, 400.0, "PUBLICATIONS", 10.0),
        (50.0, 414.0, "Rivera J. Streaming joins at scale in modern data platforms. 2022.", 10.0),
    ]
    return full + left + right + bottom


def two_column_pdf(out: Path) -> Path:
    return _pdf_lines(out / "two_column.pdf", [_two_column_items()])


def sidebar_pdf(out: Path) -> Path:
    top = [
        (50.0, 50.0, PERSON, 16.0),
        (50.0, 70.0, "PROFESSIONAL SUMMARY", 10.0),
        (50.0, 84.0, "Data engineer with ten years building analytics platforms for retail.", 10.0),
    ]
    side = [
        (40.0, 130.0, "SKILLS", 10.0), (40.0, 144.0, "Python", 10.0), (40.0, 158.0, "Spark", 10.0),
        (40.0, 186.0, "LANGUAGES", 10.0), (40.0, 200.0, "Spanish", 10.0),
    ]
    main = [(200.0, 130.0, "EXPERIENCE", 10.0)]
    y = 130.0
    for company, title, dates, bullets in TWO_COL_ROLES:
        y += 16
        main.append((200.0, y, f"{title} | {company} | {dates}", 10.0))
        for b in bullets:
            y += 14
            main.append((210.0, y, f"- {b}", 10.0))
    bottom = [
        (50.0, 300.0, "CERTIFICATIONS", 10.0),
        (50.0, 314.0, "AWS Certified Data Engineer - Associate", 10.0),
    ]
    return _pdf_lines(out / "sidebar.pdf", [top + side + main + bottom])


def _rasterize_into(doc: "fitz.Document", source: Path, dpi: int = 150) -> None:
    with fitz.open(source) as src:
        pix = src[0].get_pixmap(dpi=dpi)
    page = doc.new_page(width=612, height=792)
    page.insert_image(page.rect, pixmap=pix)


def scanned_pdf(out: Path) -> Path:
    src = single_column_pdf(out)
    doc = fitz.open()
    _rasterize_into(doc, src)
    path = out / "scanned.pdf"
    doc.save(path)
    doc.close()
    return path


def mixed_pdf(out: Path) -> Path:
    """Page 1 has native text; page 2 is an image-only scan."""
    native = _pdf_lines(out / "_mixed_native.pdf", [_two_column_items()])
    edu_only = _pdf_lines(
        out / "_mixed_scan_src.pdf",
        [[(50.0, 60.0, "CERTIFICATIONS", 12.0), (50.0, 80.0, "Certified Scrum Master 2020", 12.0)]],
    )
    doc = fitz.open(native)
    _rasterize_into(doc, edu_only)
    path = out / "mixed.pdf"
    doc.save(path)
    doc.close()
    return path


def docx_tables(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph("Professional Experience")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Data Engineer, Globex"
    t.cell(0, 1).text = "Mar 2019 - Present"
    merged = t.cell(1, 0).merge(t.cell(1, 1))
    merged.text = "Built streaming ETL pipelines for logistics data."
    d.add_paragraph("Education")
    d.add_paragraph("B.S. Physics, Tech University, 2014")
    path = out / "docx_tables.docx"
    d.save(path)
    return path


def docx_header_contact(out: Path) -> Path:
    d = Document()
    d.sections[0].header.paragraphs[0].text = f"{PERSON} | {EMAIL} | {PHONE}"
    d.add_paragraph("Experience")
    d.add_paragraph("Analyst, Initech, Jan 2015 - Feb 2019")
    d.add_paragraph("Automated weekly finance reports for leadership.", style="List Bullet")
    path = out / "docx_header_contact.docx"
    d.save(path)
    return path


def promotions_docx(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph(EMAIL)
    d.add_paragraph("Experience")
    for company, title, dates, bullets in PROMOTION_ROLES:
        d.add_paragraph(f"{title}, {company}, {dates}")
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    path = out / "promotions.docx"
    d.save(path)
    return path


def overlapping_docx(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph(EMAIL)
    d.add_paragraph("Experience")
    for (company, title, dates), bullet in zip(EXPECTED["overlapping"]["roles"], EXPECTED["overlapping"]["facts"]):
        d.add_paragraph(f"{title}, {company}, {dates}")
        d.add_paragraph(bullet, style="List Bullet")
    path = out / "overlapping.docx"
    d.save(path)
    return path


def odd_headings_docx(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph(EMAIL)
    d.add_paragraph("WHERE I'VE WORKED")
    d.add_paragraph("Data Engineer, Globex, Mar 2019 - Present")
    d.add_paragraph("Built streaming ETL pipelines for logistics data.", style="List Bullet")
    d.add_paragraph("Tools I Use:")
    d.add_paragraph("Kubernetes, Terraform")
    path = out / "odd_headings.docx"
    d.save(path)
    return path


def long_resume_docx(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph(EMAIL)
    d.add_paragraph("Experience")
    for i, (company, title, dates, n) in enumerate(LONG_ROLES):
        d.add_paragraph(f"{title}, {company}, {dates}")
        for j in range(n):
            d.add_paragraph(_long_bullet(i, j), style="List Bullet")
    path = out / "long_resume.docx"
    d.save(path)
    return path


def unknown_sections_docx(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph(EMAIL)
    d.add_paragraph("Experience")
    d.add_paragraph("Analyst, Initech, Jan 2015 - Feb 2019")
    d.add_paragraph("Automated weekly finance reports for leadership.", style="List Bullet")
    for title, line in [
        ("Publications", "Rivera J. Streaming joins at scale. 2022."),
        ("Volunteering", "Mentor, Code Club, 2020 - Present"),
        ("Patents", "US 0,000,001 Event replay system"),
    ]:
        d.add_paragraph(title)
        d.add_paragraph(line)
    path = out / "unknown_sections.docx"
    d.save(path)
    return path


def references_docx(out: Path) -> Path:
    d = Document()
    d.add_paragraph(PERSON)
    d.add_paragraph(f"{EMAIL} | {PHONE}")
    d.add_paragraph("Experience")
    d.add_paragraph("Analyst, Initech, Jan 2015 - Feb 2019")
    d.add_paragraph("Automated weekly finance reports for leadership.", style="List Bullet")
    d.add_paragraph("References")
    d.add_paragraph("Pat Manager, pat.manager@example.org, (555) 099-8877")
    path = out / "references.docx"
    d.save(path)
    return path


def template_with_sample(out: Path) -> Path:
    d = Document()
    d.sections[0].footer.paragraphs[0].text = "Client Staffing Ltd | Confidential"
    d.add_paragraph(TEMPLATE_SAMPLE_LINES[0])
    d.add_paragraph("Professional Summary")
    d.add_paragraph(TEMPLATE_SAMPLE_LINES[1])
    d.add_paragraph("Work Experience")
    d.add_paragraph(TEMPLATE_SAMPLE_LINES[2])
    path = out / "template_with_sample.docx"
    d.save(path)
    return path


BUILDERS = {
    "single_column_pdf": single_column_pdf,
    "two_column_pdf": two_column_pdf,
    "sidebar_pdf": sidebar_pdf,
    "scanned_pdf": scanned_pdf,
    "mixed_pdf": mixed_pdf,
    "docx_tables": docx_tables,
    "docx_header_contact": docx_header_contact,
    "promotions": promotions_docx,
    "overlapping": overlapping_docx,
    "odd_headings": odd_headings_docx,
    "long_resume": long_resume_docx,
    "references": references_docx,
    "unknown_sections": unknown_sections_docx,
    "template_with_sample": template_with_sample,
}


def build_all(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    return {key: builder(out) for key, builder in BUILDERS.items()}
