"""Uploaded formats keep paragraph alignment, especially a centered candidate name."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.docx_skeleton_service import render_from_docx_skeleton
from app.services.format_extraction_service import FormatExtractionService
from app.services.format_validator import document_section_types, has_critical_findings, validate_format_document
from app.services.resume_generation_service import ResumeGenerationService


def _centered_template(path) -> None:
    template = Document()
    name = template.add_paragraph("Sample Candidate")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.runs[0].bold = True
    heading = template.add_paragraph("Professional Experience")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    template.add_paragraph("Old Employer — Sample role that must not be copied")
    template.save(path)


DOC = {
    "header": {"name": "Jordan Rivera"},
    "sections": [
        {
            "type": "experience",
            "title": "Professional Experience",
            "content": [
                {
                    "company": "Acme",
                    "title": "Engineer",
                    "duration": "2020 - 2023",
                    "description": ["Shipped the billing API."],
                }
            ],
        }
    ],
}


def test_docx_format_keeps_centered_name_and_candidate_facts(tmp_path):
    path = tmp_path / "client_format.docx"
    _centered_template(path)
    schema = FormatExtractionService().extract("client_format.docx", path.read_bytes())

    assert schema["layout"]["name_alignment"] == "center"
    assert schema["layout"]["heading_alignment"] == "left"

    rendered = render_from_docx_skeleton(path, DOC, schema)
    doc = Document(io.BytesIO(rendered))
    name_para = next(p for p in doc.paragraphs if p.text == "Jordan Rivera")
    assert name_para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    heading = next(p for p in doc.paragraphs if p.text == "Professional Experience")
    assert heading.alignment in {None, WD_ALIGN_PARAGRAPH.LEFT}

    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Acme" in body
    assert "Shipped the billing API." in body
    assert "Sample Candidate" not in body
    assert "Old Employer" not in body


def test_schema_renderer_centers_name_from_saved_layout():
    service = ResumeGenerationService()
    metadata = {
        "source_type": "pdf",
        "styling": {"font_family": "Arial", "font_size_body": 10, "font_size_name": 10, "font_size_header": 10},
        "layout": {"name_alignment": "center", "heading_alignment": "left", "type": "single_column"},
        "sections": ["experience"],
        "section_order": ["experience"],
    }
    rendered = service._render_docx(DOC, metadata, logos=[])
    doc = Document(io.BytesIO(rendered))
    name_para = next(p for p in doc.paragraphs if p.text == "Jordan Rivera")
    assert name_para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "Shipped the billing API." in "\n".join(p.text for p in doc.paragraphs)


def test_left_aligned_name_stays_left(tmp_path):
    template = Document()
    name = template.add_paragraph("Left Name")
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    template.add_paragraph("Education")
    path = tmp_path / "left.docx"
    template.save(path)
    schema = FormatExtractionService().extract("left.docx", path.read_bytes())
    assert schema["layout"]["name_alignment"] == "left"


def test_summary_of_projects_heading_stays_a_projects_section():
    document = {
        "header": {"name": "Jordan Rivera"},
        "sections": [
            {"type": "bullets", "title": "Professional Summary", "content": ["Ten years building finance systems."]},
            {"type": "skills", "title": "Skills Profile", "content": ["Python"]},
            {
                "type": "experience",
                "title": "Professional Experience",
                "content": [
                    {
                        "company": "Acme",
                        "title": "Engineer",
                        "duration": "2020 - 2024",
                        "description": ["Shipped the billing API."],
                    }
                ],
            },
            {
                "type": "projects",
                "title": "Summary of Projects",
                "content": [{"title": "Billing", "description": ["Shipped the billing API."]}],
            },
            {
                "type": "education",
                "title": "Education",
                "content": [{"degree": "BS", "institution": "State University"}],
            },
        ],
    }
    spec = {
        "section_order": ["summary", "skills", "experience", "education"],
        "completeness_contract": ["summary", "skills", "experience", "education"],
    }

    assert document_section_types(document) == [
        "summary",
        "skills",
        "experience",
        "projects",
        "education",
    ]
    findings = validate_format_document(document, spec)
    assert not has_critical_findings(findings)


def test_template_header_sample_person_is_replaced(tmp_path):
    template = Document()
    header = template.sections[0].header
    header.paragraphs[0].text = "K. Naresh Kumar"
    header.add_paragraph("Java Backend developer")
    template.sections[0].footer.paragraphs[0].text = "Aptino Inc."
    template.add_paragraph("Professional Summary")
    template.add_paragraph("Sample summary that must not be copied")
    path = tmp_path / "header_format.docx"
    template.save(path)

    rendered = render_from_docx_skeleton(
        path,
        {**DOC, "header": {"name": "Jordan Rivera", "role": "Data Engineer"}},
        {"source_type": "docx", "styling": {"font_family": "Arial", "font_size_body": 10}},
    )
    doc = Document(io.BytesIO(rendered))
    header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    footer_text = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "Jordan Rivera" in header_text
    assert "Data Engineer" in header_text
    assert "K. Naresh Kumar" not in header_text
    assert "Java Backend developer" not in header_text
    assert "Aptino Inc." in footer_text


def test_credential_without_a_school_is_kept():
    service = ResumeGenerationService()
    items = service._filter_grounded_education(
        [{"degree": "Bachelor of Science", "institution": "State University"}],
        [
            {"degree": "Bachelor of Science", "institution": "State University"},
            {"degree": "CPA Certificate", "institution": ""},
        ],
    )
    assert "CPA Certificate" in [item.get("degree") for item in items]
