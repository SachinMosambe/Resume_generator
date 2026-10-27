import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml

from app.services.docx_reader import DocxLimitError, check_docx_archive, read_docx_blocks
from app.services.source_blocks import blocks_to_text


def test_tables_stay_in_document_order(fixtures):
    text = blocks_to_text(read_docx_blocks(str(fixtures["docx_tables"])))
    assert text.index("Globex") < text.index("Education")


def test_merged_cells_not_duplicated(fixtures):
    text = blocks_to_text(read_docx_blocks(str(fixtures["docx_tables"])))
    assert text.count("Built streaming ETL pipelines for logistics data.") == 1


def test_header_contact_is_extracted_first(fixtures):
    blocks = read_docx_blocks(str(fixtures["docx_header_contact"]))
    assert blocks[0].origin == "header"
    assert "jordan.rivera@example.com" in blocks[0].text


def test_footer_kept(fixtures):
    blocks = read_docx_blocks(str(fixtures["template_with_sample"]))
    assert blocks[-1].origin == "footer"


def test_list_bullets_flagged(fixtures):
    blocks = read_docx_blocks(str(fixtures["docx_header_contact"]))
    assert any(b.is_list and "finance reports" in b.text for b in blocks)
    assert "• Automated weekly finance reports" in blocks_to_text(blocks)


def test_nested_table_and_hyperlink(tmp_path):
    d = Document()
    outer = d.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "Projects"
    inner = cell.add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "Ledger rewrite"
    inner.cell(0, 1).text = "2022"
    p = d.add_paragraph("Profile: ")
    r_id = p.part.relate_to(
        "https://github.com/jrivera",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    p._p.append(
        parse_xml(
            f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}">'
            f"<w:r><w:t>GitHub</w:t></w:r></w:hyperlink>"
        )
    )
    path = tmp_path / "nested.docx"
    d.save(path)
    text = blocks_to_text(read_docx_blocks(str(path)))
    assert "Ledger rewrite | 2022" in text
    assert text.index("Projects") < text.index("Ledger rewrite")
    assert "https://github.com/jrivera" in text


def test_text_box_content_extracted_once(tmp_path):
    d = Document()
    p = d.add_paragraph("Anchor")
    p._p.append(
        parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            "<mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:txbx><w:txbxContent>"
            "<w:p><w:r><w:t>Sidebar skill Kubernetes</w:t></w:r></w:p>"
            "</w:txbxContent></wps:txbx></w:drawing></mc:Choice>"
            "<mc:Fallback><w:pict><v:textbox><w:txbxContent>"
            "<w:p><w:r><w:t>Sidebar skill Kubernetes</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
        )
    )
    path = tmp_path / "textbox.docx"
    d.save(path)
    blocks = read_docx_blocks(str(path))
    text = blocks_to_text(blocks)
    assert text.count("Sidebar skill Kubernetes") == 1
    assert any(b.origin == "textbox" for b in blocks)


def test_archive_limits(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (110 * 1024 * 1024))
    with pytest.raises(DocxLimitError):
        check_docx_archive(str(path))
