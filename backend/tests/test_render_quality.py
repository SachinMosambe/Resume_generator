import io

from docx import Document
from docx.shared import Inches

from app.services.docx_skeleton_service import render_from_docx_skeleton
from tests.fixtures.build_fixtures import TEMPLATE_SAMPLE_LINES

DOC = {
    "header": {"name": "Jordan Rivera"},
    "sections": [
        {
            "type": "experience",
            "title": "Work Experience",
            "content": [
                {"company": "Acme", "title": "Engineer", "duration": "2019 - 2021",
                 "description": ["Built the billing service in Python."]},
                {"company": "Globex", "title": "Analyst", "duration": "2016 - 2019", "description": []},
            ],
        },
        {"type": "bullets", "title": "Publications", "content": ["Streaming joins. 2022."]},
    ],
}


def _render(path):
    return Document(io.BytesIO(render_from_docx_skeleton(path, DOC, {})))


def test_headings_and_role_headers_keep_with_next(fixtures):
    doc = _render(fixtures["template_with_sample"])
    heading = next(p for p in doc.paragraphs if p.text == "Work Experience")
    assert heading.paragraph_format.keep_with_next is True
    role_header = next(p for p in doc.paragraphs if p.text.startswith("Acme"))
    assert role_header.paragraph_format.keep_with_next is True


def test_template_footer_kept_and_sample_text_removed(fixtures):
    doc = _render(fixtures["template_with_sample"])
    assert "Client Staffing Ltd" in doc.sections[0].footer.paragraphs[0].text
    body = "\n".join(p.text for p in doc.paragraphs)
    assert not any(line in body for line in TEMPLATE_SAMPLE_LINES)


def test_no_trailing_empty_paragraphs(fixtures):
    doc = _render(fixtures["template_with_sample"])
    assert doc.paragraphs[-1].text.strip() != ""


def test_leading_body_logo_is_preserved(tmp_path):
    import pymupdf as fitz

    png = tmp_path / "logo.png"
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    pix.save(str(png))
    template = Document()
    template.add_paragraph().add_run().add_picture(str(png), width=Inches(0.5))
    template.add_paragraph("Sample Candidate Name")
    path = tmp_path / "logo_template.docx"
    template.save(path)

    doc = _render(path)
    assert len(doc.inline_shapes) == 1
    assert "Sample Candidate Name" not in "\n".join(p.text for p in doc.paragraphs)
