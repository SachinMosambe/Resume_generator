import io

from docx import Document

from app.services.docx_skeleton_service import render_from_docx_skeleton
from app.services.resume_generation_service import ResumeGenerationService
from tests.fixtures.build_fixtures import PROMOTION_ROLES


def _items():
    return [{"company": c, "title": t, "duration": d, "description": list(b)} for c, t, d, b in PROMOTION_ROLES]


def test_dedupe_keeps_promotion_without_bullets():
    cleaned = ResumeGenerationService()._dedupe_experience_items(_items())
    assert [i["title"] for i in cleaned] == ["Software Engineer", "Senior Software Engineer", "Staff Engineer"]


def test_dedupe_still_drops_placeholder_stub():
    items = _items() + [{"company": "Acme Corp", "title": "Role", "duration": "", "description": []}]
    cleaned = ResumeGenerationService()._dedupe_experience_items(items)
    assert len(cleaned) == 3


def test_dedupe_collapses_same_dates_stub_into_rich_row():
    items = [
        {"company": "Acme Corp", "title": "", "duration": "Jan 2016 - Dec 2018", "description": []},
        {"company": "Acme Corp", "title": "Software Engineer", "duration": "Jan 2016 - Dec 2018",
         "description": ["Built the billing service in Python for retail clients."]},
    ]
    cleaned = ResumeGenerationService()._dedupe_experience_items(items)
    assert len(cleaned) == 1 and cleaned[0]["title"] == "Software Engineer"


def test_skeleton_renders_bulletless_role(fixtures):
    document = {
        "header": {"name": "Jordan Rivera"},
        "sections": [{"type": "experience", "title": "Work Experience", "content": _items()}],
    }
    data = render_from_docx_skeleton(fixtures["template_with_sample"], document, {})
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            text += "\n" + "\n".join(c.text for c in row.cells)
    assert "Senior Software Engineer" in text
