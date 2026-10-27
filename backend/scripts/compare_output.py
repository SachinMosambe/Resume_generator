"""Compare generated-DOCX coverage of real resumes between this checkout and another.

Usage (from backend/):
    python scripts/compare_output.py ..\\..\\other_checkout\\backend

Coverage = share of source text lines (4+ words) whose normalized text appears in the
generated Aptino-format DOCX. Prints counts only; no resume content.
"""

from __future__ import annotations

import io
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parents[1]
RESUMES = sorted(
    p
    for p in (HERE.parent / "test").rglob("*")
    if p.suffix.lower() in {".pdf", ".docx"} and "output" not in str(p).lower()
)

_DUMP = r"""
import io, json, logging, os, sys, tempfile
os.environ["UPLOAD_DIR"] = tempfile.mkdtemp()
os.environ["RESUME_LLM_CONDENSE"] = "false"
import structlog
structlog.configure(wrapper_class=structlog.make_filtering_bound_logger(logging.CRITICAL))
from docx import Document
from app.models.candidate import Candidate
from app.services.aptino_template import build_aptino_client_format
from app.services.pdf_parser import extract_text_from_document
from app.services.resume_generation_service import ResumeGenerationService
out = {}
for path in json.loads(sys.argv[1]):
    try:
        text = extract_text_from_document(path) or ""
        cand = Candidate(name="Candidate", resume_path=path, extracted_data={"raw_text": text, "resume_text": text})
        data = ResumeGenerationService().generate(cand, build_aptino_client_format("Client"))
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for r in t.rows:
                parts.extend(c.text for c in r.cells)
        out[path] = "\n".join(parts)
    except Exception as exc:
        out[path] = "ERROR " + type(exc).__name__
json.dump(out, open(sys.argv[2], "w", encoding="utf-8"))
"""


def norm(text: str) -> str:
    text = re.sub(r"[\u2010-\u2015]", "-", str(text or "")).lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def dump(backend: Path, python: str) -> dict[str, str]:
    target = Path(tempfile.mkdtemp()) / "out.json"
    paths = json.dumps([str(p) for p in RESUMES])
    subprocess.run(
        [python, "-W", "ignore", "-c", _DUMP, paths, str(target)],
        cwd=str(backend),
        check=True,
        capture_output=True,
        env={**os.environ, "PYTHONPATH": str(backend)},
    )
    return json.loads(target.read_text(encoding="utf-8"))


def main() -> None:
    other = Path(sys.argv[1]).resolve()
    sys.path.insert(0, str(HERE))
    from app.services.pdf_parser import extract_text_from_document

    this_out = dump(HERE, sys.executable)
    other_out = dump(other, sys.executable)
    print(f"{'resume':30} {'lines':>6} {'this':>6} {'other':>6}")
    for path in RESUMES:
        try:
            source = extract_text_from_document(str(path)) or ""
        except Exception as exc:  # noqa: BLE001
            print(f"{path.name[:30]:30} source error {type(exc).__name__}")
            continue
        needles = {norm(l) for l in source.splitlines() if len(l.split()) >= 4}
        needles.discard("")
        this_n = norm(this_out.get(str(path), ""))
        other_n = norm(other_out.get(str(path), ""))
        hit_this = sum(1 for n in needles if n in this_n)
        hit_other = sum(1 for n in needles if n in other_n)
        print(f"{path.name[:30]:30} {len(needles):>6} {hit_this:>6} {hit_other:>6}")


if __name__ == "__main__":
    main()
