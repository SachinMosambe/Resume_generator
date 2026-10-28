"""Read DOCX content in document order.

Order: page headers, then body paragraphs and tables as they appear (nested tables
and text boxes included), then page footers. Header/footer text is kept; the
client "name only" policy is applied at render time, not here.
"""
from __future__ import annotations

import zipfile
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.source_blocks import SourceBlock

MAX_DOCX_ENTRIES = 5_000
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_BLOCKS = 20_000

_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_PLACEHOLDER_TITLES = {"", "role", "n/a", "na", "none"}


class DocxLimitError(ValueError):
    """The DOCX archive exceeds safe size/entry limits."""


def check_docx_archive(path: str) -> None:
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_DOCX_ENTRIES:
            raise DocxLimitError("DOCX has too many archive entries")
        if sum(info.file_size for info in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise DocxLimitError("DOCX expands beyond the size limit")


class _Reader:
    def __init__(self) -> None:
        self.blocks: list[SourceBlock] = []
        self._n = 0

    def _add(self, text: str, origin: str, **kwargs: Any) -> None:
        if not text.strip():
            return
        if len(self.blocks) >= MAX_BLOCKS:
            raise DocxLimitError("DOCX has too many text blocks")
        self._n += 1
        self.blocks.append(SourceBlock(f"b{self._n:05d}", text, origin, **kwargs))

    def paragraph(self, p: Paragraph, origin: str) -> None:
        style = None
        try:
            style = p.style.name if p.style is not None else None
        except Exception:
            style = None
        p_pr = p._p.pPr
        is_list = (p_pr is not None and p_pr.numPr is not None) or bool(style and "list" in style.lower())
        runs = [r for r in p.runs if r.text.strip()]
        bold = bool(runs) and all(bool(r.bold) for r in runs)
        sizes = [r.font.size.pt for r in runs if r.font.size is not None]
        links = []
        for link in getattr(p, "hyperlinks", []) or []:
            url = getattr(link, "url", "") or ""
            if url and not url.startswith("#"):
                links.append(url)
        self._add(
            p.text or "",
            origin,
            style=style,
            is_list=is_list,
            bold=bold,
            font_size=max(sizes) if sizes else None,
            links=links,
        )
        self._textboxes(p, origin)

    def _textboxes(self, p: Paragraph, origin: str) -> None:
        for box in p._p.iter(qn("w:txbxContent")):
            if _has_ancestor(box, _MC_FALLBACK) or _has_ancestor(box, qn("w:txbxContent"), stop=p._p):
                continue
            for child in box.iterchildren():
                if child.tag == qn("w:p"):
                    self.paragraph(Paragraph(child, p), "textbox" if origin == "body" else origin)
                elif child.tag == qn("w:tbl"):
                    self.table(Table(child, p), "textbox" if origin == "body" else origin)

    def table(self, table: Table, origin: str) -> None:
        seen: set[int] = set()
        # Hold the elements. id() of a cell element is reused after it is freed,
        # which made a later merged cell look like a duplicate and dropped its text.
        pinned: list[Any] = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                element = cell._tc
                pinned.append(element)
                key = id(element)
                if key in seen:
                    continue
                seen.add(key)
                cells.append(cell)
            multi_para = any(
                sum(1 for p in cell.paragraphs if p.text.strip()) > 1 or cell.tables for cell in cells
            )
            if multi_para:
                for cell in cells:
                    self.container(cell, origin)
                continue
            texts = [cell.text.strip() for cell in cells if cell.text.strip()]
            if texts:
                links = [
                    getattr(link, "url", "")
                    for cell in cells
                    for p in cell.paragraphs
                    for link in (getattr(p, "hyperlinks", []) or [])
                    if getattr(link, "url", "") and not getattr(link, "url", "").startswith("#")
                ]
                self._add(" | ".join(texts), "table" if origin == "body" else origin, links=links)

    def container(self, container: Any, origin: str) -> None:
        for item in container.iter_inner_content():
            if isinstance(item, Paragraph):
                self.paragraph(item, origin)
            elif isinstance(item, Table):
                self.table(item, origin)


def _has_ancestor(el: Any, tag: str, stop: Any = None) -> bool:
    parent = el.getparent()
    while parent is not None and parent is not stop:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _header_footer_parts(document: Any, kind: str) -> list[Any]:
    parts = []
    seen: set[int] = set()
    for section in document.sections:
        names = (f"first_page_{kind}", kind, f"even_page_{kind}")
        for name in names:
            part = getattr(section, name, None)
            if part is None or part.is_linked_to_previous:
                continue
            key = id(part.part)
            if key in seen:
                continue
            seen.add(key)
            parts.append(part)
    return parts


def read_docx_blocks(path: str) -> list[SourceBlock]:
    check_docx_archive(path)
    document = Document(path)
    reader = _Reader()
    for header in _header_footer_parts(document, "header"):
        reader.container(header, "header")
    reader.container(document, "body")
    for footer in _header_footer_parts(document, "footer"):
        reader.container(footer, "footer")
    return reader.blocks
