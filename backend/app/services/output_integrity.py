"""Mechanical source-vs-output check for generated resumes.

Compares normalized text presence only. A "passed" status means every checked source
field was found in the document, not that the document is semantically verified.
Warnings are recruiter-facing and never contain resume content.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

from docx import Document

# Below this share of a core field found in the output, the document is marked failed.
_FAIL_RATIO = 0.5
_CORE_CHECKS = ("companies", "titles", "dates", "bullets")


@dataclass
class IntegrityReport:
    status: str
    checks: dict[str, dict] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _norm(text: Any) -> str:
    text = re.sub(r"[\u2010-\u2015]", "-", str(text or "")).lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    for section in doc.sections:
        for part in (section.header, section.footer):
            parts.extend(p.text for p in part.paragraphs)
    return _norm("\n".join(parts))


def _check(name: str, expected: list[Any], hay: str, checks: dict, warnings: list[str], label: str) -> None:
    items = list(dict.fromkeys(n for n in (_norm(x) for x in expected) if n))
    if not items:
        return
    missing = [x for x in items if x not in hay]
    checks[name] = {"ok": not missing, "expected": len(items), "found": len(items) - len(missing)}
    if missing:
        warnings.append(f"{len(missing)} of {len(items)} {label} from the source were not found in the output.")


def check_output_integrity(
    store: dict,
    docx_bytes: bytes,
    *,
    template_sample_text: str = "",
    allowed_text: str = "",
    suppressed_sections: list[str] | None = None,
    extraction_meta: dict | None = None,
) -> IntegrityReport:
    try:
        hay = _docx_text(docx_bytes)
    except Exception:
        return IntegrityReport("failed", {}, ["The generated document could not be opened."])

    checks: dict[str, dict] = {}
    warnings: list[str] = []
    name = (store.get("header") or {}).get("name")
    _check("name", [name] if name else [], hay, checks, warnings, "candidate name")

    roles = [r for r in store.get("experience") or [] if isinstance(r, dict)]
    _check("companies", [r.get("company") for r in roles], hay, checks, warnings, "employers")
    _check("titles", [r.get("title") for r in roles], hay, checks, warnings, "job titles")
    _check("dates", [r.get("duration") for r in roles], hay, checks, warnings, "employment dates")
    bullets = [b for r in roles for b in (r.get("description") or []) if isinstance(b, str)]
    _check("bullets", bullets, hay, checks, warnings, "experience bullets")

    edu = [e for e in store.get("education") or [] if isinstance(e, dict)]
    edu_fields = [e.get(k) for e in edu for k in ("degree", "institution", "school", "year")]
    _check("education", edu_fields, hay, checks, warnings, "education fields")

    grouped = store.get("skills_by_category") or {}
    skills = [s for vals in grouped.values() for s in (vals or [])] if isinstance(grouped, dict) else []
    _check("skills", skills or store.get("skills") or [], hay, checks, warnings, "skills")
    for key, label in (
        ("certifications", "certifications"),
        ("achievements", "achievements"),
        ("languages", "languages"),
    ):
        _check(key, store.get(key) or [], hay, checks, warnings, label)

    suppressed = set(suppressed_sections or [])
    extras = [
        s for s in store.get("additional_sections") or []
        if isinstance(s, dict) and s.get("title") not in suppressed
    ]
    _check("additional_sections", [l for s in extras for l in s.get("lines") or []], hay, checks, warnings,
           "lines from other sections")

    allowed = _norm(allowed_text)
    leaks = [
        n for n in (_norm(l) for l in (template_sample_text or "").splitlines())
        if len(n) >= 8 and n in hay and n not in allowed
    ]
    checks["template_leak"] = {"ok": not leaks, "expected": 0, "found": len(leaks)}
    if leaks:
        warnings.append(f"{len(leaks)} line(s) from the client template's sample content appear in the output.")

    meta = extraction_meta or {}
    unread = list(meta.get("ocr_missing_pages") or [])
    ocr = list(meta.get("ocr_pages") or [])
    checks["source_pages"] = {"ok": not unread and not ocr, "expected": len(unread) + len(ocr), "found": len(ocr)}
    if unread:
        warnings.append(
            f"{len(unread)} page(s) of the source look like scanned images and could not be read "
            "(OCR is not available); their content is missing."
        )
    if ocr:
        warnings.append(f"{len(ocr)} page(s) were read with OCR; check names, dates, and numbers.")

    for title in suppressed:
        warnings.append(f"Section '{title}' was omitted by client policy.")

    if all(c["ok"] for c in checks.values()):
        status = "passed"
    elif any(
        checks[k]["expected"] >= 2 and checks[k]["found"] < _FAIL_RATIO * checks[k]["expected"]
        for k in _CORE_CHECKS
        if k in checks
    ):
        status = "failed"
    else:
        status = "needs_review"
    return IntegrityReport(status, checks, warnings)
